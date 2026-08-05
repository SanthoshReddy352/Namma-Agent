from pptx import Presentation
from pptx.util import Emu

p = Presentation(r"C:\Users\santh\Desktop\Namma_Agent_Pitch_Deck.pptx")
print("slides:", len(p.slides))
for i, s in enumerate(p.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip().replace("\n", " | ")
            if t:
                texts.append(t)
        elif sh.shape_type == 19:  # table
            texts.append("[TABLE]")
    joined = " || ".join(texts)
    print(f"\n--- Slide {i} ---")
    print(joined[:600])